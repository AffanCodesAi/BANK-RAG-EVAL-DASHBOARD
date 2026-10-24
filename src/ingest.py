
import hashlib
import os
import re
import sys
import time
from pathlib import Path

import chromadb
import pymupdf
import torch
from docx import Document
from dotenv import load_dotenv
from sentence_transformers import SentenceTransformer


ROOT = Path(__file__).resolve().parent.parent
DOCUMENTS_DIR = ROOT / "data" / "documents"

load_dotenv(ROOT / ".env")

MODEL_NAME = os.getenv(
    "EMBEDDING_MODEL",
    "sentence-transformers/all-MiniLM-L6-v2",
)
CHROMA_PATH = Path(os.getenv("CHROMA_PATH", "./chroma_db"))
if not CHROMA_PATH.is_absolute():
    CHROMA_PATH = ROOT / CHROMA_PATH

COLLECTION_NAME = os.getenv("COLLECTION_NAME", "bank_knowledge")

BATCH_SIZE = 64
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 200


def clean_text(text):
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_text(text):
    """Split text into overlapping character-based chunks."""
    text = clean_text(text)
    chunks = []
    start = 0

    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))

        if end < len(text):
            boundary = text.rfind(" ", start + CHUNK_SIZE // 2, end)
            if boundary > start:
                end = boundary

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = max(end - CHUNK_OVERLAP, start + 1)

    return chunks


def extract_pdf(path):
    """Return (text, page_number) pairs. Page numbers start at 1."""
    results = []

    with pymupdf.open(path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = clean_text(page.get_text("text"))
            if text:
                results.append((text, page_number))

    return results


def extract_docx(path):
    """Extract paragraphs, tables, and external hyperlink URLs."""
    document = Document(path)
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                clean_text(cell.text)
                for cell in row.cells
            ]
            if any(cells):
                parts.append(" | ".join(cells))

    # Preserve hyperlink destinations, including links whose URLs
    # are not visible in paragraph text.
    links = []
    for relation in document.part.rels.values():
        if relation.reltype.endswith("/hyperlink") and relation.is_external:
            links.append(str(relation.target_ref))

    if links:
        parts.append("Document links:\n" + "\n".join(sorted(set(links))))

    text = clean_text("\n".join(parts))
    return [(text, 0)] if text else []


def extract_txt(path):
    text = path.read_text(encoding="utf-8", errors="replace")
    text = clean_text(text)
    return [(text, 0)] if text else []


def extract_file(path):
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".txt":
        return extract_txt(path)

    return []


def sha256_text(value):
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def main():
    if not DOCUMENTS_DIR.is_dir():
        print(f"ERROR: Documents folder not found: {DOCUMENTS_DIR}")
        sys.exit(1)

    files = sorted(
        path for path in DOCUMENTS_DIR.rglob("*")
        if path.is_file()
        and path.suffix.lower() in {".pdf", ".docx", ".txt"}
    )

    if not files:
        print(f"No PDF, DOCX, or TXT files found in {DOCUMENTS_DIR}")
        sys.exit(1)

    device = "cuda" if torch.cuda.is_available() else "cpu"
    print(f"Documents found: {len(files)}")
    print(f"Embedding device: {device}")
    if device == "cuda":
        print(f"GPU: {torch.cuda.get_device_name(0)}")

    print("Loading embedding model...")
    start = time.perf_counter()
    model = SentenceTransformer(MODEL_NAME, device=device)
    print(f"Model loaded in {time.perf_counter() - start:.2f}s")

    client = chromadb.PersistentClient(path=str(CHROMA_PATH))
    collection = client.get_or_create_collection(
        name=COLLECTION_NAME,
        metadata={"hnsw:space": "cosine"},
    )

    total_chunks = 0
    skipped_files = 0

    for file_number, path in enumerate(files, start=1):
        source = path.relative_to(ROOT).as_posix()
        print(f"\n[{file_number}/{len(files)}] {path.name}")

        try:
            extracted_pages = extract_file(path)
        except Exception as exc:
            print(f"  ERROR extracting file: {exc}")
            skipped_files += 1
            continue

        chunks = []
        for page_text, page_number in extracted_pages:
            for chunk in split_text(page_text):
                chunks.append((chunk, page_number))

        if not chunks:
            print("  WARNING: No extractable text; skipping.")
            print("  If this is a scanned PDF, it may need OCR.")
            skipped_files += 1
            continue

        file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        ids = []
        documents = []
        metadatas = []

        for chunk_index, (chunk, page_number) in enumerate(chunks):
            chunk_id = sha256_text(
                f"{source}|{file_hash}|{page_number}|{chunk_index}|{chunk}"
            )

            ids.append(chunk_id)
            documents.append(chunk)
            metadatas.append({
                "source": source,
                "file_name": path.name,
                "page": page_number,
                "chunk_index": chunk_index,
                "file_sha256": file_hash,
            })

        try:
            # Embed in batches on the selected device.
            for start_index in range(0, len(documents), BATCH_SIZE):
                end_index = start_index + BATCH_SIZE

                batch_docs = documents[start_index:end_index]
                batch_ids = ids[start_index:end_index]
                batch_metadata = metadatas[start_index:end_index]

                embeddings = model.encode(
                    batch_docs,
                    batch_size=BATCH_SIZE,
                    show_progress_bar=False,
                    convert_to_numpy=True,
                    normalize_embeddings=True,
                )

                collection.upsert(
                    ids=batch_ids,
                    documents=batch_docs,
                    embeddings=embeddings.tolist(),
                    metadatas=batch_metadata,
                )

            # Remove outdated chunks for this source only after
            # all current chunks have been successfully upserted.
            existing = collection.get(
                where={"source": source},
                include=["metadatas"],
            )
            current_ids = set(ids)
            stale_ids = [
                existing_id
                for existing_id in existing["ids"]
                if existing_id not in current_ids
            ]

            if stale_ids:
                collection.delete(ids=stale_ids)

            total_chunks += len(chunks)
            print(f"  Chunks stored: {len(chunks)}")

        except Exception as exc:
            print(f"  ERROR storing chunks: {exc}")
            print("  Re-run ingestion after resolving the error.")
            skipped_files += 1

    print("\n========== INGESTION SUMMARY ==========")
    print(f"Files discovered: {len(files)}")
    print(f"Files skipped or failed: {skipped_files}")
    print(f"Chunks processed successfully: {total_chunks}")
    print(f"Collection: {COLLECTION_NAME}")
    print(f"Collection total: {collection.count()}")
    print(f"Database: {CHROMA_PATH}")
    print(f"Embedding device: {device}")


if __name__ == "__main__":
    main()
